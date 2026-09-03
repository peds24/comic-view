"""Import a Comic Geeks Excel export of physical comics into the library.

No network calls here — same "local pass first, network pass in enrich"
split as the digital scanner. A physical comic that matches an existing
digital record (same series + issue number) is merged into that record by
adding "physical" to its formats; everything else becomes a new
physical-only record.
"""
from __future__ import annotations

import hashlib
import re

import openpyxl

from library.models import ComicRecord

SHEET_NAME = "Comics"

_YEAR_RANGE_SUFFIX_RE = re.compile(r"\s*\(\d{4}\s*-\s*(?:\d{4}|present)\)\s*$", re.IGNORECASE)
_ISSUE_RE = re.compile(r"#(\d+)")

# Titles containing any of these are collected editions / variants — never
# attempted against a single-issue digital match, always physical-only.
_NON_MATCHABLE_KEYWORDS = (
    "annual", "hc", "tp", "omnibus", "compendium", "special",
    "printing", "variant", "edition", "collection",
)


def load_rows(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SHEET_NAME]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def normalize_series(raw: str) -> str:
    return _YEAR_RANGE_SUFFIX_RE.sub("", raw or "").strip()


def extract_issue(full_title: str) -> str | None:
    match = _ISSUE_RE.search(full_title or "")
    return match.group(1) if match else None


def normalize_issue(issue: str) -> str:
    if issue.isdigit():
        return str(int(issue))
    return issue


def is_matchable(full_title: str, issue_number: str | None) -> bool:
    if issue_number is None:
        return False
    lowered = (full_title or "").lower()
    return not any(keyword in lowered for keyword in _NON_MATCHABLE_KEYWORDS)


def parse_row(row: dict) -> dict | None:
    """Returns a partial ComicRecord-shaped dict, or None if the row should be skipped."""
    if row.get("In Collection") != 1:
        return None

    full_title = row.get("Full Title") or ""
    series = normalize_series(row.get("Series Name") or "")
    issue_number = extract_issue(full_title)

    year = None
    release_date = row.get("Release Date")
    if release_date and str(release_date)[:4].isdigit():
        year = int(str(release_date)[:4])

    return {
        "title": full_title,
        "series": series or None,
        "issue_number": issue_number,
        "publisher": row.get("Publisher Name") or None,
        "year": year,
        "status": "read" if row.get("Marked Read") == 1 else "unread",
        "matchable": is_matchable(full_title, issue_number),
    }


def compute_physical_id(series: str, full_title: str) -> str:
    digest = hashlib.sha1(f"{series.lower()}|{full_title.lower()}".encode()).hexdigest()
    return f"xlsx-{digest[:16]}"


def find_digital_match(
    records: dict[str, ComicRecord], series: str, issue_number: str
) -> ComicRecord | None:
    target_series = (series or "").lower()
    target_issue = normalize_issue(issue_number)
    for record in records.values():
        if "digital" not in record.formats:
            continue
        if (record.series or "").lower() != target_series:
            continue
        if record.issue_number is None:
            continue
        if normalize_issue(record.issue_number) != target_issue:
            continue
        return record
    return None


def import_physical(records: dict[str, ComicRecord], rows: list[dict]) -> tuple[int, int, int]:
    """Merges/adds physical entries into records (mutated in place).

    Returns (merged_count, new_count, skipped_count).
    """
    merged = 0
    new = 0
    skipped = 0

    for row in rows:
        parsed = parse_row(row)
        if parsed is None:
            skipped += 1
            continue

        match = None
        if parsed["matchable"] and parsed["series"] and parsed["issue_number"]:
            match = find_digital_match(records, parsed["series"], parsed["issue_number"])

        if match is not None:
            if "physical" not in match.formats:
                match.formats.append("physical")
                merged += 1
            continue

        comic_id = compute_physical_id(parsed["series"] or "", parsed["title"])
        if comic_id in records:
            continue

        records[comic_id] = ComicRecord(
            id=comic_id,
            title=parsed["title"],
            type="comic",
            series=parsed["series"],
            issue_number=parsed["issue_number"],
            publisher=parsed["publisher"],
            year=parsed["year"],
            status=parsed["status"],
            formats=["physical"],
        )
        new += 1

    return merged, new, skipped
